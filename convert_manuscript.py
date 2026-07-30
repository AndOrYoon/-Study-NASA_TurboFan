import sys
sys.stdout.reconfigure(encoding='utf-8')

import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MD_PATH = r"C:\BMAD_PY313\Manuscript\Full-Text_Manuscript\manuscript_full_text.md"
OUT_PATH = r"C:\BMAD_PY313\Manuscript\Full-Text_Manuscript\manuscript_draft(RESS_Submission).docx"

# ── helpers ──────────────────────────────────────────────────────────────────

def set_font(run, size=11, bold=False, italic=False, color=None):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run('─' * 80)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)

def apply_inline(paragraph, text):
    """Parse bold/italic inline markers and add runs."""
    # patterns: **bold**, *italic*, `code`
    pattern = re.compile(r'(\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`)')
    pos = 0
    for m in pattern.finditer(text):
        # plain text before match
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        full = m.group(0)
        if full.startswith('**'):
            r = paragraph.add_run(m.group(2))
            r.bold = True
        elif full.startswith('*'):
            r = paragraph.add_run(m.group(3))
            r.italic = True
        elif full.startswith('`'):
            r = paragraph.add_run(m.group(4))
            r.font.name = 'Courier New'
            r.font.size = Pt(9)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])

def parse_table(doc, lines):
    """lines: list of raw '| ... |' strings (header + separator + rows)."""
    rows = []
    for line in lines:
        if re.match(r'^\s*\|[-:| ]+\|\s*$', line):
            continue  # separator row
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        rows.append(cells)
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = 'Table Grid'
    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            if ci >= ncols:
                break
            cell = table.cell(ri, ci)
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            apply_inline(p, cell_text)
            for run in p.runs:
                run.font.size = Pt(10)
            if ri == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()  # spacing after table

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(10)

# ── main conversion ───────────────────────────────────────────────────────────

def convert():
    with open(MD_PATH, encoding='utf-8') as f:
        raw = f.read()

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(3.0)
        section.right_margin  = Cm(2.5)

    # Default paragraph font
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)

    lines = raw.splitlines()
    i = 0
    in_code = False
    code_lines = []

    while i < len(lines):
        line = lines[i]

        # ── code block ──
        if line.strip().startswith('```'):
            if not in_code:
                in_code = True
                code_lines = []
                i += 1
                continue
            else:
                in_code = False
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.left_indent = Inches(0.4)
                r = p.add_run('\n'.join(code_lines))
                r.font.name = 'Courier New'
                r.font.size = Pt(9)
                doc.add_paragraph()
                i += 1
                continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # ── horizontal rule ──
        if re.match(r'^---+\s*$', line):
            add_horizontal_rule(doc)
            i += 1
            continue

        # ── H1 title ──
        if line.startswith('# ') and not line.startswith('## '):
            p = doc.add_heading(line[2:].strip(), level=1)
            p.paragraph_format.space_after = Pt(12)
            i += 1
            continue

        # ── H2 ──
        if line.startswith('## ') and not line.startswith('### '):
            p = doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue

        # ── H3 ──
        if line.startswith('### ') and not line.startswith('#### '):
            p = doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue

        # ── H4 ──
        if line.startswith('#### '):
            p = doc.add_heading(line[5:].strip(), level=4)
            i += 1
            continue

        # ── blockquote (metadata block at top) ──
        if line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(line[2:].strip().replace('**', ''))
            r.italic = True
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        # ── table ──
        if line.strip().startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i])
                i += 1
            parse_table(doc, table_lines)
            continue

        # ── figure placeholder / caption ──
        if re.match(r'^\[FIGURE', line.strip()):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            r = p.add_run(f'[{line.strip()[1:-1]}]')
            r.bold = True
            r.font.color.rgb = RGBColor(0x99, 0x00, 0x00)
            i += 1
            # next line may be caption starting with **Fig.
            if i < len(lines) and lines[i].strip().startswith('**Fig.'):
                add_caption(doc, lines[i].strip().replace('**', ''))
                i += 1
            continue

        # ── math display block $$ ... $$ ──
        if line.strip().startswith('$$'):
            math_lines = [line.strip()]
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('$$'):
                math_lines.append(lines[i].strip())
                i += 1
            if i < len(lines):
                math_lines.append(lines[i].strip())
                i += 1
            math_text = ' '.join(l for l in math_lines if l != '$$')
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(math_text)
            r.font.name = 'Cambria Math'
            r.font.size = Pt(11)
            continue

        # ── blank line ──
        if not line.strip():
            # only add paragraph spacing, not an extra blank paragraph for every blank line
            i += 1
            continue

        # ── normal paragraph ──
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        apply_inline(p, line.strip())
        for run in p.runs:
            if not run.font.name or run.font.name == 'Calibri':
                run.font.name = 'Times New Roman'
            if not run.font.size:
                run.font.size = Pt(11)
        i += 1

    doc.save(OUT_PATH)
    print(f"Saved: {OUT_PATH}")

if __name__ == '__main__':
    convert()
