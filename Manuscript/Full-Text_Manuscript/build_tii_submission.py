#!/usr/bin/env python3
"""
Build manuscript_draft(TII Submission).docx
- Base: downloaded IEEE Transactions Word template (ieee_template_base.docx)
- Content: manuscript_full_text.md
- Double-blind: author info replaced with [Anonymous]
- Styles: Title / Authors / Abstract / IndexTerms / Text / Heading 1-3 / References
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEMPLATE  = Path(__file__).parent / "ieee_template_base.docx"
MD_PATH   = Path(__file__).parent / "manuscript_full_text.md"
OUT_PATH  = Path(__file__).parent / "manuscript_draft(TII Submission).docx"

# ── inline formatting ────────────────────────────────────────────────────────
INLINE = re.compile(
    r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\$\$[^$]+\$\$|\$[^$\n]+\$)'
)

def add_inline(para, text, base_size=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            r = para.add_run(part[2:-2]); r.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            r = para.add_run(part[1:-1]); r.italic = True
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            r = para.add_run(part[1:-1])
            r.font.name = 'Courier New'
            if base_size:
                r.font.size = Pt(base_size - 1)
        elif part.startswith('$') and part.endswith('$'):
            inner = part.strip('$')
            r = para.add_run(f'[{inner}]')
            r.italic = True
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        else:
            para.add_run(part)

# ── table helper ─────────────────────────────────────────────────────────────
def add_shading(cell, fill='D9E1F2'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  fill)
    tcPr.append(shd)

def process_table(doc, lines):
    rows = []
    for ln in lines:
        if re.match(r'^\|[-:\s|]+\|$', ln):
            continue
        cells = [c.strip() for c in ln.strip('|').split('|')]
        rows.append(cells)
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=ncols)
    tbl.style = 'Normal Table'
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            txt = row[ci] if ci < len(row) else ''
            cell = tbl.cell(ri, ci)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after  = Pt(1)
            if ri == 0:
                add_shading(cell)
                r = p.add_run(txt); r.bold = True; r.font.size = Pt(8)
            else:
                add_inline(p, txt, base_size=8)
                for r in p.runs:
                    r.font.size = Pt(8)
    doc.add_paragraph()

# ── code block ───────────────────────────────────────────────────────────────
def add_code(doc, code_text):
    para = doc.add_paragraph(style='Body Text Indent')
    para.paragraph_format.space_before = Pt(3)
    para.paragraph_format.space_after  = Pt(3)
    r = para.add_run(code_text)
    r.font.name = 'Courier New'
    r.font.size = Pt(8)

# ── math block ───────────────────────────────────────────────────────────────
def add_math(doc, math_text):
    try:
        para = doc.add_paragraph(style='Equation')
    except Exception:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = para.add_run(f'[{math_text}]')
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

# ── main build ───────────────────────────────────────────────────────────────
def build():
    doc = Document(str(TEMPLATE))

    # Clear all existing body paragraphs from the template
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)
    for t in list(doc.tables):
        t._element.getparent().remove(t._element)

    # ── Title block (double-blind) ──
    title_para = doc.add_paragraph(style='Title')
    title_para.add_run(
        "From Fleet Normalization to Fault-Mode Gating: "
        "A Cross-Dataset Ablation Study of Turbofan "
        "Remaining Useful Life Prediction"
    )

    authors_para = doc.add_paragraph(style='Authors')
    r = authors_para.add_run("[Anonymous Authors — Double-Blind Submission]")
    r.italic = True
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # ── Abstract ──
    abs_label = doc.add_paragraph(style='Abstract')
    abs_label.add_run(
        "Abstract—Turbofan remaining useful life (RUL) prediction pipelines "
        "embed several interdependent design choices — RUL clipping threshold, "
        "sensor normalization, model architecture, and training loss — yet their "
        "individual contributions are rarely isolated or evaluated across all NASA CMAPSS "
        "operating conditions simultaneously. We present a controlled ablation study on all "
        "four sub-datasets (FD001–FD004), covering (i) five RUL clipping "
        "thresholds, (ii) seven normalization strategies including RevIN, "
        "(iii) four fault-mode architectures, and (iv) seven loss functions "
        "crossed with all clipping values, totalling over 800 LSTM training runs (five "
        "random seeds each) with all comparisons subject to Wilcoxon rank-sum tests "
        "corrected by Benjamini-Hochberg FDR. Fleet-level min-max normalization "
        "outperforms all per-unit and instance-level alternatives on FD001, FD002, and "
        "FD004, but exhibits anomalously high inter-seed variance on FD003 "
        "(RMSE std = 12.86 vs. ≤1.84 elsewhere) — a variance "
        "signature traced to co-existing HPC and fan fault modes. An unsupervised "
        "attention-gate model (M3) resolves this by routing engines to fault-specific "
        "branches using only the first five observed flight cycles, reducing FD003 RMSE "
        "by 65.8 % (14.78 ± 1.32 vs. 43.23 ± 0.18) "
        "and NASA Score by 98.8 %, while remaining immune to the test-time cluster "
        "collapse that degrades GMM hard-routing by 75.4 % on multi-condition FD004. "
        "No custom loss function outperforms MSE after multiple-comparison correction "
        "(96 BH-corrected tests); removing RUL clipping inflates NASA prognostic scores "
        "by up to six orders of magnitude regardless of loss choice. These results establish "
        "a three-tier design hierarchy — label engineering, fault-mode architecture, "
        "loss function — providing actionable, evidence-based guidance for practitioners "
        "deploying LSTM-based RUL predictors in industrial maintenance systems."
    )

    # ── Index Terms ──
    try:
        kw = doc.add_paragraph(style='IndexTerms')
    except Exception:
        kw = doc.add_paragraph()
    kw.add_run("Index Terms—")
    r2 = kw.add_run(
        "Remaining useful life prediction, prognostics and health management, "
        "turbofan engine, CMAPSS, LSTM, fault-mode gating, normalization ablation, "
        "loss function, Benjamini-Hochberg correction."
    )
    r2.italic = True

    # ── Parse & render body ──
    lines = MD_PATH.read_text(encoding='utf-8').split('\n')
    i = 0
    first_h1_seen = False   # skip the title line in MD (already added above)
    in_abstract   = False   # skip abstract block (already added above)
    in_references = False
    ref_head_done = False

    while i < len(lines):
        line     = lines[i]
        stripped = line.strip()

        # blank
        if not stripped:
            i += 1
            continue
        # blockquotes → skip
        if stripped.startswith('>'):
            i += 1
            continue
        # horizontal rule → skip
        if re.match(r'^-{3,}$', stripped):
            i += 1
            continue

        # H1
        if re.match(r'^# ', stripped):
            title_text = re.sub(r'^# ', '', stripped)
            if not first_h1_seen:
                first_h1_seen = True   # already added title above
                i += 1
                continue
            # Section heading
            in_abstract = False
            if title_text.strip().lower().startswith('abstract'):
                in_abstract = True
                i += 1
                continue
            p = doc.add_paragraph(style='Heading 1')
            p.add_run(title_text)
            i += 1
            continue

        # H2
        if re.match(r'^## ', stripped):
            h2_text = re.sub(r'^## ', '', stripped)
            if h2_text.strip().lower() == 'abstract':
                in_abstract = True
                i += 1
                continue
            if h2_text.strip().lower() == 'references':
                in_references = True
                if not ref_head_done:
                    try:
                        rh = doc.add_paragraph(style='Reference Head')
                    except Exception:
                        rh = doc.add_paragraph(style='Heading 2')
                    rh.add_run('References')
                    ref_head_done = True
                i += 1
                continue
            in_abstract = False
            p = doc.add_paragraph(style='Heading 1')
            p.add_run(h2_text)
            i += 1
            continue

        # H3
        if re.match(r'^### ', stripped):
            h3_text = re.sub(r'^### ', '', stripped)
            p = doc.add_paragraph(style='Heading 2')
            p.add_run(h3_text)
            i += 1
            continue

        # H4
        if re.match(r'^#### ', stripped):
            h4_text = re.sub(r'^#### ', '', stripped)
            p = doc.add_paragraph(style='Heading 3')
            r = p.add_run(h4_text); r.italic = True
            i += 1
            continue

        # skip abstract body (already rendered)
        if in_abstract:
            i += 1
            continue

        # fenced code block
        if stripped.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1
            add_code(doc, '\n'.join(code_lines))
            continue

        # display math $$…$$
        if stripped == '$$':
            math_lines = []
            i += 1
            while i < len(lines) and lines[i].strip() != '$$':
                math_lines.append(lines[i].strip())
                i += 1
            i += 1
            add_math(doc, '  '.join(math_lines))
            continue
        if stripped.startswith('$$') and stripped.endswith('$$') and len(stripped) > 4:
            add_math(doc, stripped[2:-2].strip())
            i += 1
            continue

        # table
        if stripped.startswith('|'):
            tbl_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                tbl_lines.append(lines[i].strip())
                i += 1
            process_table(doc, tbl_lines)
            continue

        # bullet list
        if re.match(r'^[-*]\s+', stripped):
            p = doc.add_paragraph(style='List Bullet')
            add_inline(p, stripped[2:])
            for r in p.runs: r.font.size = Pt(9)
            i += 1
            continue

        # numbered list
        nm = re.match(r'^\d+\.\s+(.*)', stripped)
        if nm:
            p = doc.add_paragraph(style='List Number')
            add_inline(p, nm.group(1))
            for r in p.runs: r.font.size = Pt(9)
            i += 1
            continue

        # references
        if in_references and re.match(r'^\[\d+\]', stripped):
            try:
                p = doc.add_paragraph(style='References')
            except Exception:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
            add_inline(p, stripped)
            for r in p.runs: r.font.size = Pt(8)
            i += 1
            continue

        # double-blind filter: skip author-identifying lines
        if re.match(r'^\*{0,2}(Authors?:|Affiliation:|Correspondence:|Email:|Author:|ORCID:)', stripped, re.I):
            i += 1
            continue

        # figure placeholder: [FIGURE N: Title]
        m_fig = re.match(r'^\[FIGURE (\d+): (.+)\]$', stripped)
        if m_fig:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after  = Pt(2)
            r = p.add_run(f'[Figure {m_fig.group(1)}: {m_fig.group(2)}]')
            r.italic = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
            i += 1
            continue

        # figure caption: **Fig. N.** ...
        if re.match(r'^\*\*Fig\. \d', stripped):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after  = Pt(12)
            add_inline(p, stripped)
            for r in p.runs:
                r.font.size = Pt(9)
            i += 1
            continue

        # normal body paragraph
        try:
            p = doc.add_paragraph(style='Text')
        except Exception:
            p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1

    doc.save(str(OUT_PATH))
    print(f"Done → {OUT_PATH}")

if __name__ == '__main__':
    build()



