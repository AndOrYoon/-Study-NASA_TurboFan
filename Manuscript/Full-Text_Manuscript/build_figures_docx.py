#!/usr/bin/env python3
"""Build Figures.docx — all 8 manuscript figures with captions embedded."""

from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

FIGURES_DIR = Path(__file__).parent.parent / "Figures"
OUT_PATH    = Path(__file__).parent / "Figures.docx"

FIGURES = [
    {
        "num": 1,
        "file": "Fig1_degradation_trends.png",
        "title": "Sensor Degradation Trajectories",
        "caption": (
            "Illustrative sensor degradation trajectories for representative engines in FD003 "
            "(single operating condition, two fault modes). Selected sensors exhibiting high RUL "
            "correlation (s2, s3, s4, s7, s11, s12) show visually distinct degradation patterns "
            "between HPC-fault and fan-fault engines, motivating both fault-mode-aware modeling (H6) "
            "and fleet-level normalization that preserves inter-engine degradation contrast (H5)."
        ),
    },
    {
        "num": 2,
        "file": "Fig2_rul_clipping.png",
        "title": "Piecewise Linear RUL Label and Clipping",
        "caption": (
            "Piecewise linear RUL labeling scheme with threshold-based clipping. The raw RUL "
            "decreases linearly from the maximum cycle but is clipped at threshold c to account "
            "for the healthy phase in which degradation is undetectable. The shaded region "
            "illustrates the effect of varying c ∈ {75, 100, 125, 130, ∞}. Over-clipping (small c) "
            "truncates degradation information for long-lifetime engines; no clipping (c = ∞) "
            "permits unbounded targets that destabilize model training and the NASA prognostic score."
        ),
    },
    {
        "num": 3,
        "file": "Fig3_H2_clipping_rmse.png",
        "title": "Effect of RUL Clipping on Prediction Accuracy",
        "caption": (
            "RMSE heatmap across four CMAPSS sub-datasets (FD001–FD004) and five clipping "
            "thresholds (mean over 20 runs, Ridge regression baseline). Each cell reports the mean "
            "RMSE over 20 independent runs. The clip = 125 configuration achieves the lowest or "
            "near-lowest RMSE in all datasets, confirming the empirical optimality of the "
            "industry-standard threshold. The absence of clipping (clip = None) produces catastrophic "
            "predictions on FD003 (RMSE = 56.1 cycles), where unbounded RUL targets amplify "
            "residuals near the beginning of engine life."
        ),
    },
    {
        "num": 4,
        "file": "Fig4_H5_normalization_rmse.png",
        "title": "RMSE Comparison Across Normalization Strategies",
        "caption": (
            "Mean RMSE heatmap (5 seeds) for seven normalization strategies (N1–N7) across four "
            "CMAPSS sub-datasets. Fleet-level min-max normalization (N1) achieves the lowest RMSE "
            "on FD001 (14.14), FD002 (14.31), and FD004 (14.60). Per-unit normalization methods "
            "(N3–N6) consistently underperform by removing between-engine degradation contrast. "
            "RevIN (N7) performs competitively on single-condition FD001 (14.92) but degrades "
            "substantially on multi-condition datasets (FD002: 18.16, FD004: 18.28). The "
            "anomalously high variance of N1 on FD003 (std = 12.86) is attributed to unresolved "
            "fault-mode mixing, addressed in H6."
        ),
    },
    {
        "num": 5,
        "file": "Fig5_H5_statistical_test.png",
        "title": "Statistical Significance of Normalization Differences (vs N1)",
        "caption": (
            "Pairwise statistical comparison of each normalization method against fleet min-max (N1) "
            "using the Wilcoxon rank-sum test with Benjamini-Hochberg FDR correction (α = 0.05). "
            "Effect sizes are reported as Cohen's d. Per-unit methods (N3–N6) and RevIN (N7) are "
            "significantly inferior on FD001, FD002, and FD004 (|d| ≥ 1.66 in all significant cases). "
            "No alternative achieves statistically significant superiority over N1. FD003 shows no "
            "significant differences due to high seed-to-seed variance from mixed fault modes."
        ),
    },
    {
        "num": 6,
        "file": "Fig6_H6_gmm_clustering.png",
        "title": "GMM Fault-Mode Cluster Quality (Phase 1)",
        "caption": (
            "Gaussian Mixture Model (GMM, K = 2) clustering results for FD003 and FD004 in Phase 1. "
            "(a) Silhouette scores across three feature variants: AB_full (all cycles), AB_slope "
            "(degradation slope features), and AB_late (final 20% of cycles). AB_late achieves the "
            "highest silhouette (FD003: 0.858; FD004: 0.855), confirming that fault-mode separation "
            "is most pronounced in late-life degradation. (b) BIC scores consistently favor K = 2, "
            "validating the two-cluster hypothesis corresponding to HPC-fault and fan-fault "
            "degradation modes. The AB_full variant (silhouette: FD003 = 0.761, FD004 = 0.750) "
            "is used in Phase 2 to retain temporal diversity across the full engine lifetime."
        ),
    },
    {
        "num": 7,
        "file": "Fig7_H6_model_comparison.png",
        "title": "Fault-Mode Separation Model Comparison (M0–M3)",
        "caption": (
            "RMSE comparison (mean ± std, 5 seeds) of four fault-mode separation architectures on "
            "FD003 and FD004. M3 (Attention Gate) achieves RMSE = 14.78 ± 1.32 on FD003, a 65.8% "
            "reduction versus M0 (43.23 ± 0.18), while remaining statistically equivalent to M0 on "
            "FD004 (28.33 ± 1.03 vs 28.05 ± 1.74). M1 collapses on FD004 (49.20 ± 7.17) due to "
            "test-time cluster assignment collapse (247:1 ratio), demonstrating the fragility of "
            "hard routing at inference time. M3's end-to-end learning circumvents this failure mode "
            "and maintains low variance across seeds."
        ),
    },
    {
        "num": 8,
        "file": "Fig8_H7_clip_loss_interaction.png",
        "title": "RUL Clipping × Loss Function Interaction (NASA Score)",
        "caption": (
            "Mean NASA prognostic score (lower is better) across all four datasets as a function "
            "of RUL clipping threshold and loss function (560 LSTM training runs, 5 seeds per "
            "configuration). RUL clipping dominates loss function choice: clip = None yields "
            "catastrophic NASA scores (up to 2.4 × 10⁶) regardless of loss function, while "
            "clip ∈ {125, 130} stabilizes training for all losses. Among clipped configurations, "
            "standard MSE (L1) with clip_130 achieves the best mean NASA score (24.39), and no "
            "custom loss function achieves a statistically significant improvement after "
            "Benjamini-Hochberg FDR correction (α = 0.05)."
        ),
    },
]


def build():
    doc = Document()

    for sec in doc.sections:
        sec.top_margin    = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin   = Inches(1.25)
        sec.right_margin  = Inches(1.25)

    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(11)

    # Cover title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Manuscript Figures")
    r.bold = True
    r.font.size = Pt(14)
    r.font.name = 'Times New Roman'

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(
        "From Fleet Normalization to Fault-Mode Gating: "
        "A Systematic Ablation Study of Turbofan Remaining Useful Life Prediction"
    )
    r2.italic = True
    r2.font.size = Pt(10)
    r2.font.name = 'Times New Roman'
    r2.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    doc.add_paragraph()

    for fig in FIGURES:
        img_path = FIGURES_DIR / fig["file"]

        # Figure number header
        p_num = doc.add_paragraph()
        p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_num.paragraph_format.space_before = Pt(18)
        r_num = p_num.add_run(f"Fig. {fig['num']}. — {fig['title']}")
        r_num.bold = True
        r_num.font.size = Pt(11)
        r_num.font.name = 'Times New Roman'

        # Image
        if img_path.exists():
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(4)
            p_img.paragraph_format.space_after  = Pt(4)
            run = p_img.add_run()
            run.add_picture(str(img_path), width=Inches(5.5))
        else:
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_miss = p_img.add_run(f"[Image not found: {img_path.name}]")
            r_miss.italic = True
            r_miss.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
            r_miss.font.size = Pt(9)

        # Caption
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after  = Pt(20)
        r_bold = p_cap.add_run(f"Fig. {fig['num']}. ")
        r_bold.bold = True
        r_bold.font.size = Pt(9)
        r_bold.font.name = 'Times New Roman'
        r_cap = p_cap.add_run(fig["caption"])
        r_cap.font.size = Pt(9)
        r_cap.font.name = 'Times New Roman'

        # Divider
        doc.add_paragraph()

    doc.save(str(OUT_PATH))
    print(f"Done → {OUT_PATH}")


if __name__ == '__main__':
    build()
